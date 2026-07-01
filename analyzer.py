import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from math import isfinite

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill


YELLOW_FILL = PatternFill(fill_type='solid', fgColor='FFF2CC')
NO_FILL = PatternFill(fill_type=None)

COL_DATE = '근태일'
COL_WEEKDAY = '요일'
COL_DEPT = '소속명'
COL_NAME = '교직원'
COL_EMPLOYEE_NO = '교직원번호'
COL_WORK_TYPE = '근무유형'
COL_WORK_SHAPE = '근무형태'
COL_CLOCK_IN = '출근시간'
COL_CLOCK_OUT = '퇴근시간'
COL_NEXT_DAY = '익일여부'
COL_CHECK = '체크'
COL_STATUS = '근무상황'
COL_RESULT = '검증결과'
COL_REASON = '검증사유'

LATE = '지참'
EARLY_LEAVE = '조퇴'
OUTING = '외출'
CHILDCARE = '육아시간'
ANNUAL_LEAVE = '연가'
SICK_LEAVE = '병가'
PUBLIC_LEAVE = '공가'
CONGRATULATORY_LEAVE = '청가'
SPECIAL_LEAVE = '특별휴가'
BUSINESS_TRIP = '출장'
ALT_DAY_OFF = '대체휴무'
LONG_SERVICE_LEAVE = '장기재직휴가'
PREGNANCY_REDUCTION = '임신기근무시간단축'
FAMILY_CARE = '가족돌봄휴가'
HALF_DAY = '반일'
HALF_DAY_ALIASES = ('반일', '반차')
MORNING_HALF_DAY = '오전'
AFTERNOON_HALF_DAY = '오후'

START_APPROVALS = {
    LATE,
    CHILDCARE,
    ANNUAL_LEAVE,
    SICK_LEAVE,
    PUBLIC_LEAVE,
    BUSINESS_TRIP,
    PREGNANCY_REDUCTION,
    FAMILY_CARE,
}
END_APPROVALS = {
    EARLY_LEAVE,
    CHILDCARE,
    ANNUAL_LEAVE,
    SICK_LEAVE,
    PUBLIC_LEAVE,
    BUSINESS_TRIP,
    PREGNANCY_REDUCTION,
    FAMILY_CARE,
}
FULL_DAY_APPROVALS = {
    ANNUAL_LEAVE,
    SICK_LEAVE,
    PUBLIC_LEAVE,
    CONGRATULATORY_LEAVE,
    SPECIAL_LEAVE,
    BUSINESS_TRIP,
    ALT_DAY_OFF,
    LONG_SERVICE_LEAVE,
    FAMILY_CARE,
}
KNOWN_TIMED_APPROVALS = START_APPROVALS | END_APPROVALS | FULL_DAY_APPROVALS | {OUTING}

DEFAULT_START = 9 * 60
DEFAULT_END = 18 * 60
LUNCH_START = 12 * 60
LUNCH_END = 13 * 60
HALF_DAY_WORK_MINUTES = 4 * 60
WORK_MINUTE_TOLERANCE = 5
AS_OF_DATE = os.environ.get('ATTENDANCE_CHECKER_AS_OF_DATE') or date.today().isoformat()
NO_APPROVAL_MARKERS = {
    '',
    '-',
    '없음',
    '해당없음',
    '해당 없음',
    '정상',
    '없슴',
    '무',
    'N/A',
}
NO_APPROVAL_MARKER_KEYS = None
TIME_EXPR_PATTERN = r'(?:\d{1,2}:\d{2}|\d{3,4}|\d{1,2}(?:\.\d+)?|\d{1,2}\s*시\s*(?:\d{1,2}\s*분?)?)'
RANGE_SEPARATOR_PATTERN = r'\s*(?:~|-|〜|∼|－|–|—)\s*'
TIME_RANGE_RE = re.compile(fr'({TIME_EXPR_PATTERN}){RANGE_SEPARATOR_PATTERN}({TIME_EXPR_PATTERN})')
APPROVAL_INTERVAL_RE = re.compile(
    fr'([^,()]+)(?:\(([^()]*)\))?\(({TIME_EXPR_PATTERN}){RANGE_SEPARATOR_PATTERN}({TIME_EXPR_PATTERN})\)'
)


@dataclass(frozen=True)
class Approval:
    kind: str
    subtype: str
    start: int | None
    end: int | None
    raw: str


@dataclass(frozen=True)
class RowDecision:
    row_idx: int
    category: str
    anomaly_type: str
    issues: tuple[str, ...]
    base_start: int | None
    base_end: int | None
    adjusted_start: int | None
    adjusted_end: int | None


def normalize_text(value):
    if value is None:
        return ''
    text = unicodedata.normalize('NFKC', str(value))
    text = re.sub(r'[\u200B-\u200D\uFEFF]', '', text)
    return text.strip()


def time_to_minutes(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, timedelta):
        return int(round(value.total_seconds() / 60)) % (24 * 60)
    if isinstance(value, (int, float)):
        numeric = float(value)
        if not isfinite(numeric) or numeric < 0:
            return None

        if numeric < 1:
            return int(round(numeric * 24 * 60)) % (24 * 60)

        if 1 <= numeric <= 24:
            return int(round(numeric * 60)) % (24 * 60)

        if numeric.is_integer() and 100 <= numeric <= 2359:
            hour = int(numeric) // 100
            minute = int(numeric) % 100
            if 0 <= hour <= 23 and 0 <= minute <= 59:
                return hour * 60 + minute

        serial_fraction = numeric % 1
        if serial_fraction:
            return int(round(serial_fraction * 24 * 60)) % (24 * 60)
        return None

    text = normalize_text(value)
    if re.fullmatch(r'\d{1,2}(?:\.\d+)?', text):
        return time_to_minutes(float(text))

    match = re.search(r'(\d{1,2}):(\d{2})', text)
    if not match:
        match = re.search(r'(\d{1,2})\s*시\s*(?:(\d{1,2})\s*분?)?', text)
    if not match and re.fullmatch(r'\d{3,4}', text):
        hour = int(text) // 100
        minute = int(text) % 100
        if 0 <= hour <= 23 and 0 <= minute <= 59:
            return hour * 60 + minute
        return None
    if not match:
        return None

    hour = int(match.group(1))
    minute = int(match.group(2) or 0)
    if 0 <= hour <= 23 and 0 <= minute <= 59:
        return hour * 60 + minute
    return None


def format_minutes(value):
    if value is None:
        return ''
    day_offset = value // (24 * 60)
    value %= 24 * 60
    text = f'{value // 60:02d}:{value % 60:02d}'
    if day_offset:
        return f'+{day_offset}일 {text}'
    return text


def format_time(value):
    return format_minutes(time_to_minutes(value))


def format_cell(value):
    if value is None:
        return ''
    formatted_time = format_time(value)
    if formatted_time:
        return formatted_time
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def format_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return normalize_text(value)


def find_header(ws):
    required = {COL_DATE, COL_WORK_TYPE, COL_CLOCK_IN, COL_CLOCK_OUT, COL_STATUS}
    for row_idx in range(1, min(ws.max_row, 20) + 1):
        values = [normalize_text(ws.cell(row_idx, col_idx).value) for col_idx in range(1, ws.max_column + 1)]
        if required.issubset(values):
            return row_idx, {name: values.index(name) + 1 for name in values if name}
    return None, {}


def row_value(ws, row_idx, header, column_name):
    col_idx = header.get(column_name)
    if col_idx is None:
        return None
    return ws.cell(row=row_idx, column=col_idx).value


def schedule_for(ws, row_idx, header):
    work_type = normalize_text(row_value(ws, row_idx, header, COL_WORK_TYPE))
    work_shape = normalize_text(row_value(ws, row_idx, header, COL_WORK_SHAPE))
    match = TIME_RANGE_RE.search(work_shape)
    if match:
        return time_to_minutes(match.group(1)), time_to_minutes(match.group(2))
    if work_type:
        return DEFAULT_START, DEFAULT_END
    return None, None


def parse_approvals(value):
    text = normalize_text(value)
    if is_no_approval_marker(text):
        return []

    approvals = []
    for match in APPROVAL_INTERVAL_RE.finditer(text):
        approvals.append(
            Approval(
                kind=normalize_text(match.group(1)),
                subtype=normalize_text(match.group(2)),
                start=time_to_minutes(match.group(3)),
                end=time_to_minutes(match.group(4)),
                raw=match.group(0),
            )
        )

    for token in [part.strip() for part in text.split(',') if part.strip()]:
        if is_no_approval_marker(token):
            continue
        range_match = TIME_RANGE_RE.search(token)
        if range_match:
            if not APPROVAL_INTERVAL_RE.search(token):
                kind, subtype = parse_kind_subtype(token[:range_match.start()])
                approvals.append(
                    Approval(
                        kind=kind,
                        subtype=subtype,
                        start=time_to_minutes(range_match.group(1)),
                        end=time_to_minutes(range_match.group(2)),
                        raw=token,
                    )
                )
            continue
        kind, subtype = parse_kind_subtype(token)
        if kind:
            approvals.append(
                Approval(
                    kind=kind,
                    subtype=subtype,
                    start=None,
                    end=None,
                    raw=token,
                )
            )
    return approvals


def parse_kind_subtype(value):
    text = normalize_text(value).strip()
    match = re.match(r'([^()]+)(?:\(([^()]*)\))?$', text)
    if not match:
        return text.strip('() '), ''
    return normalize_text(match.group(1)), normalize_text(match.group(2))


def normalize_no_space(value):
    return re.sub(r'\s+', '', normalize_text(value)).upper()


def no_approval_marker_keys():
    global NO_APPROVAL_MARKER_KEYS
    if NO_APPROVAL_MARKER_KEYS is None:
        NO_APPROVAL_MARKER_KEYS = {normalize_no_space(marker) for marker in NO_APPROVAL_MARKERS}
    return NO_APPROVAL_MARKER_KEYS


def is_no_approval_marker(value):
    return normalize_no_space(value) in no_approval_marker_keys()


def check_flags(value):
    text = normalize_no_space(value)
    if is_no_approval_marker(text):
        return set()
    if re.fullmatch(r'[ABC](?:[/,;·]?[ABC])*', text):
        return set(re.findall(r'[ABC]', text))
    return set()


def append_issue_once(issues, issue):
    if issue not in issues:
        issues.append(issue)


def has_start_side_approval(approvals, base_start, clock_in):
    if base_start is None or clock_in is None:
        return False
    for approval in approvals:
        if approval.start is None or approval.end is None:
            continue
        if not kind_in(approval.kind, START_APPROVALS):
            continue
        if approval.start <= base_start < approval.end:
            uncovered = interval_minutes(approval.end, clock_in, exclude_lunch=True) if clock_in > approval.end else 0
            if uncovered <= WORK_MINUTE_TOLERANCE:
                return True
    return False


def has_start_side_approval_candidate(approvals, base_start):
    if base_start is None:
        return False
    return any(
        approval.start is not None
        and approval.end is not None
        and kind_in(approval.kind, START_APPROVALS)
        and approval.start <= base_start < approval.end
        for approval in approvals
    )


def has_end_side_approval(approvals, base_end, clock_out):
    if base_end is None or clock_out is None:
        return False
    for approval in approvals:
        if approval.start is None or approval.end is None:
            continue
        if not kind_in(approval.kind, END_APPROVALS):
            continue
        starts_at_clock_out = approval.start <= clock_out <= approval.end
        starts_after_non_work_gap = clock_out < approval.start and interval_minutes(clock_out, approval.start, exclude_lunch=True) <= WORK_MINUTE_TOLERANCE
        if starts_at_clock_out or starts_after_non_work_gap:
            return True
    return False


def has_end_side_approval_candidate(approvals, base_end):
    if base_end is None:
        return False
    return any(
        approval.start is not None
        and approval.end is not None
        and kind_in(approval.kind, END_APPROVALS)
        and approval.start <= base_end <= approval.end
        for approval in approvals
    )


def add_work_minutes(start, work_minutes):
    current = start
    remaining = work_minutes
    while remaining > 0:
        if LUNCH_START <= current < LUNCH_END:
            current = LUNCH_END
            continue

        next_stop = LUNCH_START if current < LUNCH_START else None
        if next_stop is None:
            return current + remaining

        available = next_stop - current
        if remaining <= available:
            return current + remaining
        current = LUNCH_END
        remaining -= available
    return current


def covers_schedule(approval, start, end):
    if not kind_in(approval.kind, FULL_DAY_APPROVALS) or is_half_day_approval(approval):
        return False
    if approval.start is None and approval.end is None:
        return True
    return (
        start is not None
        and end is not None
        and approval.start is not None
        and approval.end is not None
        and approval.start <= start
        and approval.end >= end
    )


def kind_in(kind, candidates):
    return any(candidate in kind for candidate in candidates)


def interval_minutes(start, end, exclude_lunch=True):
    if start is None or end is None:
        return 0
    if end < start:
        end += 24 * 60
    total = max(0, end - start)
    if not exclude_lunch:
        return total

    lunch_start = LUNCH_START
    lunch_end = LUNCH_END
    if end > 24 * 60:
        lunch_start += 24 * 60
        lunch_end += 24 * 60
    overlap = max(0, min(end, lunch_end) - max(start, lunch_start))
    return max(0, total - overlap)


def required_work_minutes(start, end):
    return interval_minutes(start, end, exclude_lunch=True)


def iter_approval_minutes(approval):
    if approval.start is None or approval.end is None:
        return
    if not kind_in(approval.kind, KNOWN_TIMED_APPROVALS):
        return

    start = approval.start
    end = approval.end
    if end < start:
        end += 24 * 60

    counted = 0
    for minute in range(start, end):
        minute_of_day = minute % (24 * 60)
        if approval.kind != CHILDCARE and LUNCH_START <= minute_of_day < LUNCH_END:
            continue
        yield minute
        counted += 1
        if approval.kind == CHILDCARE and counted >= 2 * 60:
            break


def approved_minutes(approval):
    if is_half_day_approval(approval) and approval.start is None:
        return HALF_DAY_WORK_MINUTES
    return len(set(iter_approval_minutes(approval) or []))


def total_approved_minutes(approvals):
    timed_minutes = set()
    fixed_minutes = 0
    for approval in approvals:
        if is_half_day_approval(approval) and approval.start is None:
            fixed_minutes += HALF_DAY_WORK_MINUTES
            continue
        timed_minutes.update(iter_approval_minutes(approval) or [])
    return fixed_minutes + len(timed_minutes)


def is_recognized_approval(approval):
    if kind_in(approval.kind, KNOWN_TIMED_APPROVALS):
        return True
    if is_half_day_approval(approval):
        return True
    return False


def approval_key(approval):
    return normalize_no_space(f'{approval.kind} {approval.subtype} {approval.raw}')


def is_half_day_approval(approval):
    if not kind_in(approval.kind, FULL_DAY_APPROVALS):
        return False
    key = approval_key(approval)
    return any(normalize_no_space(alias) in key for alias in HALF_DAY_ALIASES)


def is_morning_half_day(approval):
    return is_half_day_approval(approval) and normalize_no_space(MORNING_HALF_DAY) in approval_key(approval)


def adjusted_boundaries(start, end, approvals):
    adjusted_start = start
    adjusted_end = end
    for approval in approvals:
        if approval.start is None or approval.end is None:
            continue
        if (
            adjusted_start is not None
            and kind_in(approval.kind, START_APPROVALS)
            and approval.start <= adjusted_start <= approval.end
        ):
            adjusted_start = max(adjusted_start, approval.end)
        if (
            adjusted_end is not None
            and kind_in(approval.kind, END_APPROVALS)
            and approval.start <= adjusted_end <= approval.end
        ):
            adjusted_end = min(adjusted_end, approval.start)

    for approval in approvals:
        if is_half_day_approval(approval) and approval.start is None and start is not None:
            half_day_boundary = add_work_minutes(start, HALF_DAY_WORK_MINUTES)
            if is_morning_half_day(approval):
                adjusted_start = max(adjusted_start, half_day_boundary)
            else:
                adjusted_end = min(adjusted_end, half_day_boundary)
    return adjusted_start, adjusted_end


def normalized_clock_out(clock_in, clock_out, next_day):
    if clock_out is None:
        return None
    try:
        is_next_day = int(next_day or 0) != 0
    except (TypeError, ValueError):
        is_next_day = False
    if is_next_day or (clock_in is not None and clock_out < clock_in):
        return clock_out + 24 * 60
    return clock_out


def decide_row(ws, row_idx, header):
    base_start, base_end = schedule_for(ws, row_idx, header)
    clock_in = time_to_minutes(row_value(ws, row_idx, header, COL_CLOCK_IN))
    raw_clock_out = time_to_minutes(row_value(ws, row_idx, header, COL_CLOCK_OUT))
    clock_out = normalized_clock_out(clock_in, raw_clock_out, row_value(ws, row_idx, header, COL_NEXT_DAY))
    approvals = parse_approvals(row_value(ws, row_idx, header, COL_STATUS))
    erp_flags = check_flags(row_value(ws, row_idx, header, COL_CHECK))
    adjusted_start, adjusted_end = adjusted_boundaries(base_start, base_end, approvals)
    row_date = format_date(row_value(ws, row_idx, header, COL_DATE))

    issues = []
    category = 'ok'

    if base_start is None:
        if clock_in is not None or raw_clock_out is not None:
            return RowDecision(
                row_idx,
                'review',
                '검토',
                ('비근무일 출퇴근 기록',),
                base_start,
                base_end,
                adjusted_start,
                adjusted_end,
            )
        return RowDecision(row_idx, category, '', tuple(), base_start, base_end, adjusted_start, adjusted_end)

    if any(covers_schedule(approval, base_start, base_end) for approval in approvals):
        return RowDecision(row_idx, category, '', tuple(), base_start, base_end, adjusted_start, adjusted_end)

    if row_date == AS_OF_DATE and raw_clock_out is None:
        if clock_in is not None and adjusted_start is not None and clock_in > adjusted_start:
            issues.append(f'기준일 지참 후보 {format_minutes(clock_in)}>{format_minutes(adjusted_start)}')
        issues.append('기준일 퇴근 미기록')
        return RowDecision(
            row_idx,
            'pending',
            '보류',
            tuple(issues),
            base_start,
            base_end,
            adjusted_start,
            adjusted_end,
        )

    if clock_in is None and raw_clock_out is None:
        return RowDecision(
            row_idx,
            'anomaly',
            '결근의심',
            ('근무일 출퇴근 모두 없음',),
            base_start,
            base_end,
            adjusted_start,
            adjusted_end,
        )

    if clock_in is None:
        issues.append('출근시간 누락')
    if raw_clock_out is None:
        issues.append('퇴근시간 누락')

    if clock_in is not None and raw_clock_out is not None:
        required = required_work_minutes(base_start, base_end)
        actual = interval_minutes(clock_in, clock_out, exclude_lunch=True)
        approved = total_approved_minutes(approvals)

        recognized_approvals = any(is_recognized_approval(approval) for approval in approvals)

        if approvals:
            shortage = required - actual - approved
            if shortage > WORK_MINUTE_TOLERANCE:
                issues.append(
                    f'실근무+승인시간 부족 {actual + approved}분/{required}분'
                )

        if not approvals or not recognized_approvals:
            if adjusted_start is not None and clock_in > adjusted_start:
                append_issue_once(issues, f'미승인 지참 {format_minutes(clock_in)}>{format_minutes(adjusted_start)}')
            if adjusted_end is not None and clock_out is not None and clock_out < adjusted_end:
                append_issue_once(issues, f'미승인 조퇴 {format_minutes(clock_out)}<{format_minutes(adjusted_end)}')
        else:
            if (
                ('A' in erp_flags or has_start_side_approval_candidate(approvals, base_start))
                and adjusted_start is not None
                and clock_in > adjusted_start
                and not has_start_side_approval(approvals, base_start, clock_in)
            ):
                append_issue_once(issues, f'미승인 지참 {format_minutes(clock_in)}>{format_minutes(adjusted_start)}')
            if (
                ('B' in erp_flags or has_end_side_approval_candidate(approvals, base_end))
                and adjusted_end is not None
                and clock_out is not None
                and clock_out < adjusted_end
                and not has_end_side_approval(approvals, base_end, clock_out)
            ):
                append_issue_once(issues, f'미승인 조퇴 {format_minutes(clock_out)}<{format_minutes(adjusted_end)}')

    if not issues:
        return RowDecision(row_idx, category, '', tuple(), base_start, base_end, adjusted_start, adjusted_end)

    if '근무일 출퇴근 모두 없음' in issues:
        anomaly_type = '결근의심'
    elif any('지참' in issue or '출근' in issue for issue in issues) and any('조퇴' in issue or '퇴근' in issue for issue in issues):
        anomaly_type = '지참/조퇴'
    elif any('지참' in issue or '출근' in issue for issue in issues):
        anomaly_type = '지참'
    elif any('조퇴' in issue or '퇴근' in issue for issue in issues):
        anomaly_type = '조퇴'
    else:
        anomaly_type = '이상치'
    return RowDecision(row_idx, 'anomaly', anomaly_type, tuple(issues), base_start, base_end, adjusted_start, adjusted_end)


def analyze_file(file_path):
    wb = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    try:
        anomalies, summary, employees, sheet_names = process_workbook(wb, apply_marks=False)
    finally:
        wb.close()

    return {
        'anomalies': anomalies,
        'summary': summary,
        'employees': employees,
        'sheets': sheet_names,
    }


def generate_marked_workbook(input_path, output_path):
    wb = load_workbook(input_path, keep_links=False)
    try:
        process_workbook(wb, apply_marks=True)
        wb.save(output_path)
    finally:
        wb.close()


def process_workbook(wb, apply_marks=False):
    anomalies = []
    employees = set()
    total_checked = 0
    total_skipped = 0
    missing_sheets = []
    sheet_names = []

    for ws in wb.worksheets:
        header_row, header = find_header(ws)
        if not header:
            missing_sheets.append(ws.title)
            continue

        sheet_names.append(ws.title)
        result_col = reason_col = None
        if apply_marks:
            result_col, reason_col = prepare_result_columns(ws, header_row)

        for row_idx in range(header_row + 1, ws.max_row + 1):
            if not row_has_data(ws, row_idx, header):
                continue

            name = normalize_text(row_value(ws, row_idx, header, COL_NAME))
            if name:
                employees.add(name)

            decision = decide_row(ws, row_idx, header)
            if decision.base_start is None:
                total_skipped += 1
            else:
                total_checked += 1

            if decision.category in {'anomaly', 'pending', 'review'}:
                item = build_result_item(ws, row_idx, header, decision, len(anomalies) + 1)
                anomalies.append(item)
                if apply_marks:
                    mark_row(ws, row_idx, decision, result_col, reason_col)

    summary = {
        'total_checked': total_checked,
        'total_skipped': total_skipped,
        'total_anomalies': sum(1 for item in anomalies if item['category'] == 'anomaly'),
        'late_count': sum(1 for item in anomalies if item['category'] == 'anomaly' and '지참' in item['anomaly_type']),
        'early_leave_count': sum(1 for item in anomalies if item['category'] == 'anomaly' and '조퇴' in item['anomaly_type']),
        'absent_count': sum(1 for item in anomalies if item['anomaly_type'] == '결근의심'),
        'pending_count': sum(1 for item in anomalies if item['category'] == 'pending'),
        'review_count': sum(1 for item in anomalies if item['category'] == 'review'),
        'missing_sheets': missing_sheets,
        'as_of_date': AS_OF_DATE,
    }
    return anomalies, summary, sorted(employees), sheet_names


def row_has_data(ws, row_idx, header):
    for col_idx in header.values():
        if ws.cell(row=row_idx, column=col_idx).value not in (None, ''):
            return True
    return False


def build_result_item(ws, row_idx, header, decision, number):
    work_type = normalize_text(row_value(ws, row_idx, header, COL_WORK_TYPE))
    work_shape = normalize_text(row_value(ws, row_idx, header, COL_WORK_SHAPE))
    schedule = f'{format_minutes(decision.base_start)}~{format_minutes(decision.base_end)}'
    adjusted = f'{format_minutes(decision.adjusted_start)}~{format_minutes(decision.adjusted_end)}'
    detail = '; '.join(decision.issues)

    return {
        'no': number,
        'sheet': ws.title,
        'row': row_idx,
        'category': decision.category,
        'dept': normalize_text(row_value(ws, row_idx, header, COL_DEPT)),
        'name': normalize_text(row_value(ws, row_idx, header, COL_NAME)),
        'employee_no': normalize_text(row_value(ws, row_idx, header, COL_EMPLOYEE_NO)),
        'date': format_date(row_value(ws, row_idx, header, COL_DATE)),
        'weekday': normalize_text(row_value(ws, row_idx, header, COL_WEEKDAY)),
        'work_type': work_type,
        'work_shape': work_shape,
        'scheduled': schedule,
        'adjusted_schedule': adjusted,
        'checkin': format_cell(row_value(ws, row_idx, header, COL_CLOCK_IN)),
        'checkout': format_cell(row_value(ws, row_idx, header, COL_CLOCK_OUT)),
        'erp_check': normalize_text(row_value(ws, row_idx, header, COL_CHECK)),
        'anomaly_type': decision.anomaly_type,
        'anomaly_detail': detail,
        'work_status': normalize_text(row_value(ws, row_idx, header, COL_STATUS)),
        'row_values': [format_cell(ws.cell(row=row_idx, column=col_idx).value) for col_idx in range(1, ws.max_column + 1)],
    }


def prepare_result_columns(ws, header_row):
    headers = {
        normalize_text(ws.cell(header_row, col_idx).value): col_idx
        for col_idx in range(1, ws.max_column + 1)
    }
    result_col = headers.get(COL_RESULT)
    reason_col = headers.get(COL_REASON)

    next_col = ws.max_column + 1
    if result_col is None:
        result_col = next_col
        ws.cell(header_row, result_col, COL_RESULT)
        next_col += 1
    if reason_col is None:
        reason_col = next_col
        ws.cell(header_row, reason_col, COL_REASON)

    for row_idx in range(header_row + 1, ws.max_row + 1):
        result_cell = ws.cell(row_idx, result_col)
        reason_cell = ws.cell(row_idx, reason_col)
        result_cell.value = None
        reason_cell.value = None
        result_cell.comment = None
        reason_cell.comment = None
    return result_col, reason_col


def mark_row(ws, row_idx, decision, result_col, reason_col):
    reason = '; '.join(decision.issues)
    memo = f'검증결과: {decision.anomaly_type}\n검증사유: {reason}'

    result_cell = ws.cell(row_idx, result_col, decision.anomaly_type)
    reason_cell = ws.cell(row_idx, reason_col, reason)
    result_cell.comment = Comment(memo, '근태 점검 지원 프로그램(My Attendace)')
    reason_cell.comment = Comment(memo, '근태 점검 지원 프로그램(My Attendace)')

    for col_idx in range(1, reason_col + 1):
        ws.cell(row=row_idx, column=col_idx).fill = YELLOW_FILL
