import argparse
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
DEFAULT_CONFIG = ROOT / "sample_data" / "public_attendance_sample_expected.json"
DEFAULT_SAMPLE_DIR = ROOT / "sample_data"

COL_DATE = "근태일"
COL_WEEKDAY = "요일"
COL_WORK_TYPE = "근무유형"
COL_WORK_SHAPE = "근무형태"
COL_CLOCK_IN = "출근시간"
COL_CLOCK_OUT = "퇴근시간"
COL_NEXT_DAY = "익일여부"
COL_CHECK = "체크"
COL_STATUS = "근무상황"

LATE = "지참"
EARLY_LEAVE = "조퇴"
CHILDCARE = "육아시간"
ANNUAL_LEAVE = "연가"
SICK_LEAVE = "병가"
PUBLIC_LEAVE = "공가"
CONGRATULATORY_LEAVE = "청가"
SPECIAL_LEAVE = "특별휴가"
BUSINESS_TRIP = "출장"
ALT_DAY_OFF = "대체휴무"
PREGNANCY_REDUCTION = "임신기근무시간단축"
FAMILY_CARE = "가족돌봄휴가"
HALF_DAY = "반일"

START_APPROVALS = {
    LATE,
    CHILDCARE,
    ANNUAL_LEAVE,
    SICK_LEAVE,
    PUBLIC_LEAVE,
    BUSINESS_TRIP,
    PREGNANCY_REDUCTION,
    FAMILY_CARE,
}
END_APPROVALS = {
    EARLY_LEAVE,
    CHILDCARE,
    ANNUAL_LEAVE,
    SICK_LEAVE,
    PUBLIC_LEAVE,
    BUSINESS_TRIP,
    PREGNANCY_REDUCTION,
    FAMILY_CARE,
}
FULL_DAY_APPROVALS = {
    ANNUAL_LEAVE,
    SICK_LEAVE,
    PUBLIC_LEAVE,
    CONGRATULATORY_LEAVE,
    SPECIAL_LEAVE,
    BUSINESS_TRIP,
    ALT_DAY_OFF,
}

DEFAULT_START = 9 * 60
DEFAULT_END = 18 * 60
LUNCH_START = 12 * 60
LUNCH_END = 13 * 60
HALF_DAY_WORK_MINUTES = 4 * 60


@dataclass(frozen=True)
class Approval:
    kind: str
    subtype: str
    start: int | None
    end: int | None
    raw: str


@dataclass(frozen=True)
class RowResult:
    row: int
    date_value: str
    weekday: str
    work_type: str
    work_shape: str
    clock_in: str
    clock_out: str
    check: str
    status: str
    base_start: int | None
    base_end: int | None
    adjusted_start: int | None
    adjusted_end: int | None
    issues: tuple[str, ...]
    category: str


def minutes(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, timedelta):
        return int(round(value.total_seconds() / 60))
    if isinstance(value, (int, float)):
        return int(round((float(value) % 1) * 24 * 60))
    match = re.search(r"(\d{1,2}):(\d{2})", str(value))
    if not match:
        return None
    hour = int(match.group(1))
    minute = int(match.group(2))
    if 0 <= hour <= 23 and 0 <= minute <= 59:
        return hour * 60 + minute
    return None


def fmt(value):
    if value is None:
        return ""
    day_offset = value // (24 * 60)
    value %= 24 * 60
    text = f"{value // 60:02d}:{value % 60:02d}"
    if day_offset:
        return f"+{day_offset}d {text}"
    return text


def as_text(value):
    if value is None:
        return ""
    return str(value).strip()


def as_iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return as_text(value)


def find_header(ws):
    required = {COL_DATE, COL_WORK_TYPE, COL_CLOCK_IN, COL_CLOCK_OUT, COL_STATUS}
    for row_idx in range(1, min(ws.max_row, 20) + 1):
        values = [as_text(ws.cell(row_idx, col).value) for col in range(1, ws.max_column + 1)]
        if required.issubset(values):
            return {name: values.index(name) + 1 for name in values if name}
    raise ValueError(f"시트 {ws.title!r}에서 헤더 행을 찾지 못했습니다.")


def schedule_for(row):
    work_type = as_text(row.get(COL_WORK_TYPE))
    work_shape = as_text(row.get(COL_WORK_SHAPE))
    match = re.search(r"\((\d{1,2}:\d{2})\s*[~-]\s*(\d{1,2}:\d{2})\)", work_shape)
    if match:
        return minutes(match.group(1)), minutes(match.group(2))
    if work_type:
        return DEFAULT_START, DEFAULT_END
    return None, None


def parse_approvals(status):
    text = as_text(status)
    approvals = []
    interval_pattern = re.compile(
        r"([^,()]+)(?:\(([^()]*)\))?\((\d{1,2}:\d{2})\s*[~-]\s*(\d{1,2}:\d{2})\)"
    )
    for match in interval_pattern.finditer(text):
        approvals.append(
            Approval(
                kind=match.group(1).strip(),
                subtype=(match.group(2) or "").strip(),
                start=minutes(match.group(3)),
                end=minutes(match.group(4)),
                raw=match.group(0),
            )
        )

    for token in [part.strip() for part in text.split(",") if part.strip()]:
        if re.search(r"\d{1,2}:\d{2}\s*[~-]\s*\d{1,2}:\d{2}", token):
            continue
        match = re.match(r"([^()]+)(?:\(([^()]*)\))?$", token)
        if match:
            approvals.append(
                Approval(
                    kind=match.group(1).strip(),
                    subtype=(match.group(2) or "").strip(),
                    start=None,
                    end=None,
                    raw=token,
                )
            )
    return approvals


def add_work_minutes(start, work_minutes):
    current = start
    remaining = work_minutes
    while remaining > 0:
        if LUNCH_START <= current < LUNCH_END:
            current = LUNCH_END
            continue

        next_stop = LUNCH_START if current < LUNCH_START else None
        if next_stop is None:
            return current + remaining

        available = next_stop - current
        if remaining <= available:
            return current + remaining
        current = LUNCH_END
        remaining -= available
    return current


def covers_schedule(approval, start, end):
    if approval.kind not in FULL_DAY_APPROVALS or approval.subtype == HALF_DAY:
        return False
    if approval.start is None and approval.end is None:
        return True
    return (
        start is not None
        and end is not None
        and approval.start is not None
        and approval.end is not None
        and approval.start <= start
        and approval.end >= end
    )


def adjusted_boundaries(start, end, approvals):
    adjusted_start = start
    adjusted_end = end
    for approval in approvals:
        if approval.start is None or approval.end is None:
            continue
        if (
            adjusted_start is not None
            and approval.kind in START_APPROVALS
            and approval.start <= adjusted_start <= approval.end
        ):
            adjusted_start = max(adjusted_start, approval.end)
        if (
            adjusted_end is not None
            and approval.kind in END_APPROVALS
            and approval.start <= adjusted_end <= approval.end
        ):
            adjusted_end = min(adjusted_end, approval.start)

    for approval in approvals:
        if (
            approval.kind == ANNUAL_LEAVE
            and approval.subtype == HALF_DAY
            and approval.start is None
            and start is not None
        ):
            adjusted_end = min(adjusted_end, add_work_minutes(start, HALF_DAY_WORK_MINUTES))
    return adjusted_start, adjusted_end


def normalized_clock_out(clock_in, clock_out, next_day):
    if clock_out is None:
        return None
    try:
        is_next_day = int(next_day or 0) != 0
    except (TypeError, ValueError):
        is_next_day = False
    if is_next_day or (clock_in is not None and clock_out < clock_in):
        return clock_out + 24 * 60
    return clock_out


def analyze_sample(path, as_of_date):
    import analyzer as analyzer_module

    analyzer_module.AS_OF_DATE = as_of_date
    result = analyzer_module.analyze_file(path)
    rows = []
    for item in result["anomalies"]:
        rows.append(
            RowResult(
                row=item["row"],
                date_value=item["date"],
                weekday=item["weekday"],
                work_type=item["work_type"],
                work_shape=item["work_shape"],
                clock_in=item["checkin"],
                clock_out=item["checkout"],
                check=item.get("erp_check", ""),
                status=item["work_status"],
                base_start=None,
                base_end=None,
                adjusted_start=None,
                adjusted_end=None,
                issues=tuple(item["anomaly_detail"].split("; ")) if item["anomaly_detail"] else tuple(),
                category=item["category"],
            )
        )
    return rows, result.get("summary", {})


def analyze_sheet(path, as_of_date):
    rows, _ = analyze_sample(path, as_of_date)
    return rows


def row_set(results, category):
    return {result.row for result in results if result.category == category}


def print_result_details(results, category):
    category_names = {
        "anomaly": "이상치",
        "pending": "보류",
        "review": "검토",
        "ok": "정상",
    }
    category_name = category_names.get(category, category)
    selected = [result for result in results if result.category == category]
    if not selected:
        print(f"  {category_name}: 없음")
        return
    print(f"  {category_name}:")
    for result in selected:
        print(
            "   "
            f"r{result.row} {result.date_value} {result.weekday} "
            f"{result.work_type} {result.work_shape} "
            f"출근={result.clock_in or '-'} 퇴근={result.clock_out or '-'} "
            f"기준={fmt(result.base_start)}~{fmt(result.base_end)} "
            f"보정={fmt(result.adjusted_start)}~{fmt(result.adjusted_end)} "
            f"사유={list(result.issues)}"
        )
        if result.status:
            print(f"      근무상황={result.status}")


def compare_rows(label, expected, actual):
    expected_set = set(expected)
    missing = sorted(expected_set - actual)
    extra = sorted(actual - expected_set)
    if not missing and not extra:
        return True, []
    messages = []
    if missing:
        messages.append(f"{label} 기대 행 누락: {missing}")
    if extra:
        messages.append(f"{label} 예상 밖 행: {extra}")
    return False, messages


def compare_summary(expected, actual):
    if not expected:
        return True, []

    messages = []
    for key, expected_value in expected.items():
        actual_value = actual.get(key)
        if actual_value != expected_value:
            messages.append(f"요약 {key}: 기대={expected_value}, 결과={actual_value}")
    return not messages, messages


def config_value(config, korean_key, english_key=None, default=None):
    if korean_key in config:
        return config[korean_key]
    if english_key and english_key in config:
        return config[english_key]
    return default


def sample_value(sample, korean_key, english_key=None, default=None):
    if korean_key in sample:
        return sample[korean_key]
    if english_key and english_key in sample:
        return sample[english_key]
    return default


def sample_file_name(sample):
    return (
        sample_value(sample, "파일", "file")
        or sample_value(sample, "샘플_파일", "sample_file")
    )


def config_samples(config):
    samples = config_value(config, "샘플목록", "samples", [])
    if samples:
        return samples
    if sample_file_name(config):
        return [config]
    return []


def check_regulation_documents(config, sample_dir):
    documents = config_value(config, "규정_문서", "regulation_documents", [])
    all_ok = True
    if not documents:
        print("규정 문서: 설정 없음")
        return True

    print("규정 문서:")
    for document in documents:
        file_name = sample_value(document, "파일", "file")
        purpose = sample_value(document, "용도", "purpose", "")
        path = sample_dir / file_name
        if path.exists():
            print(f"  확인: {file_name} - {purpose}")
        else:
            print(f"  누락: {file_name} - {path}")
            all_ok = False
    return all_ok


def check_half_day_rule():
    examples = [
        ("시차A형 반차", 7 * 60, 11 * 60),
        ("정상근무 반차", 9 * 60, 14 * 60),
        ("시차F형 반차", 10 * 60, 15 * 60),
    ]
    all_ok = True
    print("반차 기준:")
    for label, start, expected_end in examples:
        actual_end = add_work_minutes(start, HALF_DAY_WORK_MINUTES)
        if actual_end == expected_end:
            print(f"  확인: {label} {fmt(start)} 시작 -> {fmt(actual_end)} 퇴근 가능")
        else:
            print(f"  실패: {label} {fmt(start)} 시작 -> {fmt(actual_end)}, 기대 {fmt(expected_end)}")
            all_ok = False
    return all_ok


def check_synthetic_cases():
    import analyzer as analyzer_module

    headers = [
        analyzer_module.COL_DATE,
        analyzer_module.COL_WEEKDAY,
        analyzer_module.COL_DEPT,
        analyzer_module.COL_NAME,
        analyzer_module.COL_EMPLOYEE_NO,
        analyzer_module.COL_WORK_TYPE,
        analyzer_module.COL_WORK_SHAPE,
        analyzer_module.COL_CLOCK_IN,
        analyzer_module.COL_CLOCK_OUT,
        analyzer_module.COL_NEXT_DAY,
        analyzer_module.COL_CHECK,
        analyzer_module.COL_STATUS,
    ]
    cases = [
        {
            "이름": "신청 없이 1시간 늦게 출근하고 1시간 늦게 퇴근",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "근무상황": "",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
        {
            "이름": "근무상황 없음 표기값도 신청 없음으로 처리",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "근무상황": "없음",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
        {
            "이름": "지참 신청이 있으면 늦은 출근 정상",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "근무상황": "지참(일반)(07:00~08:00)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "지참과 조퇴 신청이 모두 실제 출퇴근을 덮으면 정상",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "09:38",
            "퇴근": "14:38",
            "체크": "A/B",
            "근무상황": "조퇴(일반)(14:30~18:00), 지참(일반)(09:00~09:40)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "지참과 조퇴 신청이 있어도 승인 사이 공백이 있으면 이상치",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "09:29",
            "퇴근": "12:00",
            "체크": "A/B",
            "근무상황": "지참(일반)(09:00~09:30), 조퇴(일반)(15:30~18:00)",
            "기대분류": "anomaly",
            "기대사유": "미승인 조퇴 12:00<15:30",
        },
        {
            "이름": "시차A형 오전 근무 후 반차",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "06:59",
            "퇴근": "11:00",
            "체크": "",
            "근무상황": "연가(반일)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "시차A형 오전반차 후 11시 출근",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "11:00",
            "퇴근": "16:00",
            "체크": "",
            "근무상황": "연가(오전반일)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "정상근무 오전반차 후 14시 출근",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "14:00",
            "퇴근": "18:00",
            "체크": "",
            "근무상황": "연가(오전반차)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "오후반차 표기도 오후 반차로 처리",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "09:00",
            "퇴근": "14:00",
            "체크": "",
            "근무상황": "연가(오후반차)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "오전반차는 전일 휴가로 보지 않음",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "",
            "퇴근": "",
            "체크": "",
            "근무상황": "연가(오전반일)",
            "기대분류": "anomaly",
            "기대사유": "근무일 출퇴근 모두 없음",
        },
        {
            "이름": "병가 반일도 전일 병가로 보지 않음",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "",
            "퇴근": "",
            "체크": "",
            "근무상황": "병가(반일)",
            "기대분류": "anomaly",
            "기대사유": "근무일 출퇴근 모두 없음",
        },
        {
            "이름": "병가 오전반차 후 14시 출근",
            "근무형태": "정상근무(09:00-18:00)",
            "출근": "14:00",
            "퇴근": "18:00",
            "체크": "",
            "근무상황": "병가(오전반차)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "외출 신청만 있고 체크 A가 남으면 지참 미신청",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "체크": "A",
            "근무상황": "외출(일반)(12:00~13:00)",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
        {
            "이름": "지참 신청만 있고 체크 B가 남으면 조퇴 미신청",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "07:00",
            "퇴근": "15:00",
            "체크": "B",
            "근무상황": "지참(일반)(07:00~08:00)",
            "기대분류": "anomaly",
            "기대사유": "미승인 조퇴 15:00<16:00",
        },
        {
            "이름": "지참 신청보다 늦게 출근하면 초과 지참",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "체크": "",
            "근무상황": "지참(일반)(07:00~07:30)",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:30",
        },
        {
            "이름": "조퇴 신청보다 일찍 퇴근하면 초과 조퇴",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "07:00",
            "퇴근": "15:00",
            "체크": "",
            "근무상황": "조퇴(일반)(15:30~16:00)",
            "기대분류": "anomaly",
            "기대사유": "미승인 조퇴 15:00<15:30",
        },
        {
            "이름": "숫자 시간도 시간으로 처리",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": 8,
            "퇴근": 17,
            "체크": "",
            "근무상황": "",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
        {
            "이름": "HHMM 숫자 시간도 시간으로 처리",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": 830,
            "퇴근": 1730,
            "체크": "",
            "근무상황": "",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:30>07:00",
        },
        {
            "이름": "근무형태 한글 시간 범위도 처리",
            "근무형태": "시차A형(7시~16시)",
            "출근": "8시",
            "퇴근": "17시",
            "체크": "",
            "근무상황": "",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
        {
            "이름": "HHMM 근무형태와 신청 시간 범위도 처리",
            "근무형태": "시차A형(0700~1600)",
            "출근": "8시",
            "퇴근": "17시",
            "체크": "",
            "근무상황": "지참(일반)(0700~0800)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "긴 대시 근무형태 시간 범위도 처리",
            "근무형태": "시차A형(07:00–16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "체크": "",
            "근무상황": "",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
        {
            "이름": "체크 N/A는 A 체크로 보지 않음",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "체크": "N/A",
            "근무상황": "외출(일반)(12:00~13:00)",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "중복 승인 시간은 한 번만 계산",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:30",
            "퇴근": "15:30",
            "체크": "",
            "근무상황": "지참(일반)(07:00~08:00), 육아시간(07:30~08:30)",
            "기대분류": "anomaly",
            "기대사유": "실근무+승인시간 부족 450분/480분",
        },
        {
            "이름": "시간 괄호 없는 지참 신청도 처리",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "체크": "",
            "근무상황": "지참(일반) 07:00~08:00",
            "기대분류": "ok",
            "기대사유": "",
        },
        {
            "이름": "알 수 없는 시간 표기는 승인으로 계산하지 않음",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "16:00",
            "체크": "",
            "근무상황": "메모(07:00~08:00)",
            "기대분류": "anomaly",
            "기대사유": "실근무+승인시간 부족 420분/480분",
        },
        {
            "이름": "알 수 없는 시간 표기는 지참 신청으로도 보지 않음",
            "근무형태": "시차A형(07:00~16:00)",
            "출근": "08:00",
            "퇴근": "17:00",
            "체크": "",
            "근무상황": "메모(07:00~08:00)",
            "기대분류": "anomaly",
            "기대사유": "미승인 지참 08:00>07:00",
        },
    ]

    print("합성 회귀검사:")
    all_ok = True
    for case in cases:
        wb = Workbook()
        ws = wb.active
        for col_idx, header in enumerate(headers, 1):
            ws.cell(1, col_idx, header)
        row = [
            "2026-06-01",
            "월",
            "테스트부서",
            "홍길동",
            "1",
            "시차출퇴근제",
            case["근무형태"],
            case["출근"],
            case["퇴근"],
            0,
            case.get("체크", ""),
            case["근무상황"],
        ]
        for col_idx, value in enumerate(row, 1):
            ws.cell(2, col_idx, value)

        _, header = analyzer_module.find_header(ws)
        decision = analyzer_module.decide_row(ws, 2, header)
        detail = "; ".join(decision.issues)
        ok = decision.category == case["기대분류"] and (
            not case["기대사유"] or case["기대사유"] in detail
        )
        if ok:
            print(f"  확인: {case['이름']}")
        else:
            print(
                f"  실패: {case['이름']} - "
                f"결과={decision.category}, 사유={detail or '-'}"
            )
            all_ok = False
    return all_ok


def run(config_path, sample_dir_override=None):
    config = json.loads(Path(config_path).read_text(encoding="utf-8"))
    sample_dir = Path(sample_dir_override or config_value(config, "샘플_폴더", "sample_dir") or DEFAULT_SAMPLE_DIR)
    as_of_date = config_value(config, "기준일", "as_of_date") or date.today().isoformat()
    samples = config_samples(config)
    all_ok = True

    print(f"설정 파일: {config_path}")
    print(f"샘플 폴더: {sample_dir}")
    print(f"기준일: {as_of_date}")
    if not check_regulation_documents(config, sample_dir):
        all_ok = False
    if not check_half_day_rule():
        all_ok = False
    if not check_synthetic_cases():
        all_ok = False

    for sample in samples:
        file_name = sample_file_name(sample)
        path = sample_dir / file_name
        print(f"\n## {file_name}")
        if not path.exists():
            print(f"  오류: 샘플 파일이 없습니다: {path}")
            all_ok = False
            continue

        results, summary = analyze_sample(path, as_of_date)
        actual_anomalies = row_set(results, "anomaly")
        actual_pending = row_set(results, "pending")
        actual_review = row_set(results, "review")

        checks = [
            compare_summary(sample_value(sample, "기대_요약", "expected_summary", {}), summary),
            compare_rows("이상치", sample_value(sample, "기대_이상치_행", "expected_anomaly_rows", []), actual_anomalies),
            compare_rows("보류", sample_value(sample, "기대_보류_행", "expected_pending_rows", []), actual_pending),
            compare_rows("검토", sample_value(sample, "기대_검토_행", "expected_review_rows", []), actual_review),
        ]
        sample_ok = all(ok for ok, _ in checks)
        all_ok = all_ok and sample_ok

        print(f"  이상치 행: {sorted(actual_anomalies)}")
        print(f"  보류 행: {sorted(actual_pending)}")
        print(f"  검토 행: {sorted(actual_review)}")
        if not sample_ok:
            for ok, messages in checks:
                if not ok:
                    for message in messages:
                        print(f"  실패: {message}")
        print_result_details(results, "anomaly")
        print_result_details(results, "pending")

    print("\n통과" if all_ok else "\n실패")
    return 0 if all_ok else 1


def main():
    parser = argparse.ArgumentParser(description="샘플 엑셀로 복무 이상치 판정 규칙을 검증합니다.")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="하네스 JSON 설정 파일 경로")
    parser.add_argument("--sample-dir", help="샘플 xlsx 파일이 있는 폴더")
    args = parser.parse_args()
    raise SystemExit(run(args.config, args.sample_dir))


if __name__ == "__main__":
    main()
